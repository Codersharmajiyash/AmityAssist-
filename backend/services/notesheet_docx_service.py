"""Notesheet DOCX generation service.

Produces a formatted Word document (.docx) containing the notesheet content,
approval trail, edit audit log, and digital signature placeholders.
"""

from __future__ import annotations

import io
from datetime import datetime, timezone
from typing import Any, Dict, List

from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT


def _add_heading_block(doc: Document, text: str, level: int = 1) -> None:
    heading = doc.add_heading(text, level=level)
    for run in heading.runs:
        run.font.color.rgb = RGBColor(0x00, 0x3B, 0x5C)


def _add_metadata_table(doc: Document, data: Dict[str, Any]) -> None:
    """Add a 2-column metadata summary table."""
    fields = [
        ("Reference No", data.get("reference_no", "—")),
        ("Title", data.get("title", "—")),
        ("Category", data.get("category", "—")),
        ("Student ID", data.get("student_id") or "N/A"),
        ("Created By", data.get("created_by", "—")),
        ("Current Stage", data.get("current_stage", "—")),
        ("Status", data.get("status", "—")),
        ("Created At", data.get("created_at", "—")),
        ("Updated At", data.get("updated_at", "—")),
    ]
    table = doc.add_table(rows=len(fields), cols=2)
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (label, value) in enumerate(fields):
        row = table.rows[i]
        row.cells[0].text = label
        row.cells[1].text = str(value) if value else "—"
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.style.font.size = Pt(10)


def _add_content_section(doc: Document, content: Any) -> None:
    """Render notesheet structured content fields."""
    _add_heading_block(doc, "Notesheet Content", level=2)
    if isinstance(content, dict):
        for key, value in content.items():
            label = key.replace("_", " ").title()
            p = doc.add_paragraph()
            run_label = p.add_run(f"{label}: ")
            run_label.bold = True
            run_label.font.size = Pt(11)
            run_value = p.add_run(str(value) if value else "—")
            run_value.font.size = Pt(11)
    elif isinstance(content, str):
        doc.add_paragraph(content)
    else:
        doc.add_paragraph(str(content) if content else "No content available.")


def _add_approval_trail(doc: Document, signatures: List[Dict[str, Any]]) -> None:
    """Render approval / signature trail as a table."""
    _add_heading_block(doc, "Approval & Signature Trail", level=2)
    if not signatures:
        doc.add_paragraph("No signatures recorded yet.")
        return

    table = doc.add_table(rows=1, cols=6)
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["#", "Officer", "Role / Stage", "Action", "Date", "Comments"]
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True
                run.font.size = Pt(9)

    for idx, sig in enumerate(signatures, 1):
        row = table.add_row()
        row.cells[0].text = str(idx)
        row.cells[1].text = str(sig.get("officer_name", "—"))
        row.cells[2].text = str(sig.get("role", sig.get("stage", "—")))
        row.cells[3].text = str(sig.get("action", "—"))
        row.cells[4].text = str(sig.get("signed_at", sig.get("timestamp", "—")))
        row.cells[5].text = str(sig.get("comments", ""))
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.style.font.size = Pt(9)


def _add_edit_audit_trail(doc: Document, edits: List[Dict[str, Any]]) -> None:
    """Render in-flight edit audit log."""
    _add_heading_block(doc, "In-Flight Edit Audit Trail", level=2)
    if not edits:
        doc.add_paragraph("No in-flight modifications recorded.")
        return

    table = doc.add_table(rows=1, cols=5)
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Field", "Old Value", "New Value", "Officer", "Reason"]
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True
                run.font.size = Pt(9)

    for edit in edits:
        row = table.add_row()
        row.cells[0].text = str(edit.get("field_name", "—"))
        row.cells[1].text = str(edit.get("old_value", "—"))
        row.cells[2].text = str(edit.get("new_value", "—"))
        row.cells[3].text = str(edit.get("officer_id", "—"))
        row.cells[4].text = str(edit.get("reason", "—"))
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.style.font.size = Pt(9)


def _add_digital_signature_block(doc: Document, data: Dict[str, Any]) -> None:
    """Add a visual digital signature verification block."""
    _add_heading_block(doc, "Digital Signature Verification", level=2)

    signatures = data.get("signatures") or []
    if not signatures:
        doc.add_paragraph("No digital signatures attached to this notesheet.")
        return

    for sig in signatures:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        officer = sig.get("officer_name", "Authorized Officer")
        role = sig.get("role", sig.get("stage", ""))
        timestamp = sig.get("signed_at", sig.get("timestamp", ""))
        action = sig.get("action", "SIGNED")
        sig_id = sig.get("officer_id", "")

        run = p.add_run(f"✓ Digitally Signed by: {officer}")
        run.bold = True
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(0x00, 0x69, 0x5C)

        details = doc.add_paragraph()
        details.add_run(f"   Role: {role}  |  Action: {action}  |  Date: {timestamp}  |  ID: {sig_id}").font.size = Pt(9)


def generate_notesheet_docx(notesheet_data: Dict[str, Any]) -> bytes:
    """Generate a complete Word document from notesheet data.

    Returns the .docx file content as bytes suitable for streaming.
    """
    doc = Document()

    # Page margins
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    # ── Document Header ─────────────────────────────────────────────
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title_p.add_run("UNIASSIST — Official Digital Notesheet")
    run.bold = True
    run.font.size = Pt(18)
    run.font.color.rgb = RGBColor(0x00, 0x3B, 0x5C)

    subtitle_p = doc.add_paragraph()
    subtitle_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ref_no = notesheet_data.get("reference_no", "")
    sub_run = subtitle_p.add_run(f"Reference: {ref_no}")
    sub_run.font.size = Pt(12)
    sub_run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    generated_p = doc.add_paragraph()
    generated_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    gen_run = generated_p.add_run(
        f"Generated: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')}"
    )
    gen_run.font.size = Pt(9)
    gen_run.font.color.rgb = RGBColor(0x99, 0x99, 0x99)

    doc.add_paragraph("")  # Spacer

    # ── Metadata ────────────────────────────────────────────────────
    _add_heading_block(doc, "Notesheet Summary", level=2)
    _add_metadata_table(doc, notesheet_data)
    doc.add_paragraph("")

    # ── Content ─────────────────────────────────────────────────────
    content = notesheet_data.get("content")
    if content:
        _add_content_section(doc, content)
        doc.add_paragraph("")

    # ── Approval Trail ──────────────────────────────────────────────
    signatures = notesheet_data.get("signatures") or []
    _add_approval_trail(doc, signatures)
    doc.add_paragraph("")

    # ── Edit Audit Trail ────────────────────────────────────────────
    edits = notesheet_data.get("edit_trail") or []
    _add_edit_audit_trail(doc, edits)
    doc.add_paragraph("")

    # ── Digital Signature Block ─────────────────────────────────────
    _add_digital_signature_block(doc, notesheet_data)

    # ── Footer ──────────────────────────────────────────────────────
    doc.add_paragraph("")
    footer_p = doc.add_paragraph()
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    f_run = footer_p.add_run(
        "This document was generated by UNIASSIST Digital Notesheet System. "
        "Verify authenticity via the reference number at the university portal."
    )
    f_run.font.size = Pt(8)
    f_run.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
    f_run.italic = True

    # Serialize to bytes
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.read()
